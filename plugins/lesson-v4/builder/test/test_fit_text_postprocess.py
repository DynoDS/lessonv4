from __future__ import annotations

import importlib.util
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "builder" / "scripts" / "fit_text_postprocess.py"
SPEC = importlib.util.spec_from_file_location("fit_text_postprocess", SCRIPT)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


def add_box(slide, name, text, x, y, w, h, size=14, bold=True):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = "Comic Sans MS"
    run.font.bold = bold
    run.font.size = Pt(size)
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    return shape


def run_sizes(shape):
    return [
        round(run.font.size.pt)
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.font.size is not None
    ]


def test_explicit_projected_floor_refuses_text_that_only_fits_at_ten_points(capsys):
    with tempfile.TemporaryDirectory() as temp_dir:
        deck = Path(temp_dir) / "small-table.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = add_box(slide, "GROWFIT__table__40__MIN20__cell", "Children learn together at school.", 0, 0, 2.5, 0.28)
        assert not MODULE.measure_shape(shape, 40, 10)["hit_floor"]
        assert MODULE.measure_shape(shape, 40, 20)["hit_floor"]
        presentation.save(deck)
        MODULE.process(str(deck), floor_pt=10)
        assert "hit 20pt floor" in capsys.readouterr().err
        assert min(run_sizes(Presentation(deck).slides[0].shapes[0])) >= 20


def test_readable_shared_table_keeps_header_and_evidence_at_same_size(capsys):
    with tempfile.TemporaryDirectory() as temp_dir:
        deck = Path(temp_dir) / "readable-table.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        add_box(slide, "GROWFIT__table__40__MIN20__header", "Material", 0, 0, 4, 0.6)
        add_box(slide, "GROWFIT__table__54__MIN20__body", "Wood and metal", 0, 0.6, 4, 1.2)
        presentation.save(deck)
        MODULE.process(str(deck), floor_pt=10)
        assert "OVERLOAD" not in capsys.readouterr().err
        shapes = Presentation(deck).slides[0].shapes
        assert run_sizes(shapes[0]) == run_sizes(shapes[1])
        assert min(run_sizes(shapes[0])) >= 20


def test_group_uses_one_largest_safe_size_without_changing_boxes():
    with tempfile.TemporaryDirectory() as temp_dir:
        deck = Path(temp_dir) / "grouped.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        group = "table-column-0-0-9000-5000"
        first = add_box(
            slide,
            f"GROWFIT__{group}__54__electrical-appliance",
            "Electrical appliance",
            0.5,
            0.5,
            3.0,
            1.3,
        )
        second = add_box(
            slide,
            f"GROWFIT__{group}__54__mains-electricity",
            "Mains electricity",
            0.5,
            2.0,
            3.0,
            1.3,
        )
        original_geometry = [
            (first.left, first.top, first.width, first.height),
            (second.left, second.top, second.width, second.height),
        ]
        presentation.save(deck)

        MODULE.process(str(deck), floor_pt=10)

        checked = Presentation(deck)
        shapes = [shape for shape in checked.slides[0].shapes if shape.has_text_frame]
        sizes = [run_sizes(shape)[0] for shape in shapes]
        assert len(set(sizes)) == 1
        assert sizes[0] > 14
        assert [
            (shape.left, shape.top, shape.width, shape.height)
            for shape in shapes
        ] == original_geometry

        one_point_larger_fails = []
        for shape in shapes:
            result = MODULE.measure_shape(shape, sizes[0] + 1, sizes[0] + 1)
            one_point_larger_fails.append(result["hit_floor"])
        assert any(one_point_larger_fails)


def test_unmarked_text_keeps_shrink_only_behaviour():
    with tempfile.TemporaryDirectory() as temp_dir:
        deck = Path(temp_dir) / "ordinary.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        add_box(slide, "ordinary", "Short", 0.5, 0.5, 5.0, 2.0, size=14)
        presentation.save(deck)

        MODULE.process(str(deck), floor_pt=10)

        checked = Presentation(deck)
        shape = next(shape for shape in checked.slides[0].shapes if shape.has_text_frame)
        assert run_sizes(shape)[0] == 14


def test_appliance_detective_reaches_the_measured_shared_maxima():
    with tempfile.TemporaryDirectory() as temp_dir:
        deck = Path(temp_dir) / "appliance-detective.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        instruction = add_box(
            slide,
            "GROWFIT__instruction__40__appliance-detective-instruction",
            "Use each photograph to name the object, decide whether it is an electrical appliance, and if it is, identify its power source. Add one piece of evidence for each decision.",
            0.897,
            0.75,
            5.32,
            1.188,
            size=12,
        )
        category_text = ["Electrical appliance", "Mains electricity", "Battery"]
        meaning_text = [
            "Made to use electricity to do its job",
            "Designed to connect to a socket",
            "Designed to use a battery as its power source",
        ]
        for index, text in enumerate(category_text):
            add_box(
                slide,
                f"GROWFIT__category-column__54__category-{index}",
                text,
                7.007,
                1.85 + index * 1.72,
                2.933,
                1.72,
                size=14,
            )
        for index, text in enumerate(meaning_text):
            add_box(
                slide,
                f"GROWFIT__meaning-column__54__meaning-{index}",
                text,
                9.94,
                1.85 + index * 1.72,
                2.933,
                1.72,
                size=14,
                bold=False,
            )
        presentation.save(deck)

        MODULE.process(str(deck), floor_pt=10)

        checked = Presentation(deck)
        by_group = {"instruction": [], "category-column": [], "meaning-column": []}
        for shape in checked.slides[0].shapes:
            directive = MODULE.grow_fit_directive(shape.name)
            if directive is not None:
                by_group[directive[0]].append(run_sizes(shape)[0])

        assert by_group["instruction"] == [16]
        assert len(set(by_group["category-column"])) == 1
        assert by_group["category-column"][0] >= 40
        assert len(set(by_group["meaning-column"])) == 1
        assert by_group["meaning-column"][0] >= 28


PROMPT_FIELDS = [
    "What is it?",
    "Is it an electrical appliance?",
    "How is it powered?",
    "What job does it do?",
]


def add_field_box(slide, name, fields, x, y, w, h, size, space_after_pt, breaks=True):
    """A card's field prompt, written the way the evidence-card helper writes it.

    `breaks=False` writes the same words as one flowing run, which is what the
    measurement used to think it was looking at either way.
    """
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    text = ("\n" if breaks else " ").join(fields)
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Comic Sans MS"
    run.font.bold = True
    run.font.size = Pt(size)
    if space_after_pt:
        paragraph.space_after = Pt(space_after_pt)
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    return shape


def test_hard_line_breaks_are_measured_as_the_lines_they_draw():
    """The Year 4 appliances deck: four field prompts, and "What job does it
    do?" printed below the card it belonged in. Every deterministic check passed
    because the measurement packed the four prompts across their own line breaks
    and counted five lines where PowerPoint drew seven."""
    with tempfile.TemporaryDirectory() as temp_dir:
        deck = Path(temp_dir) / "prompt-card.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        add_field_box(
            slide, "prompt", PROMPT_FIELDS, 0.2, 0.2, 4.083, 2.941, 32, 8
        )
        presentation.save(deck)

        MODULE.process(str(deck), floor_pt=10)

        checked = Presentation(deck)
        shape = next(s for s in checked.slides[0].shapes if s.has_text_frame)
        chosen = run_sizes(shape)[0]
        assert chosen < 32, f"the overflowing prompt was left at {chosen}pt"
        assert chosen >= 10
        # And the size it chose is the largest that really fits: one point more
        # does not.
        assert MODULE.measure_shape(shape, chosen + 1, chosen + 1)["hit_floor"]


def test_the_same_words_without_breaks_keep_their_size():
    """Discrimination: the fix must not shrink text that genuinely fits. The
    same words as one flowing line wrap to five lines and fill the box."""
    with tempfile.TemporaryDirectory() as temp_dir:
        deck = Path(temp_dir) / "flowing.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        add_field_box(
            slide,
            "flowing",
            PROMPT_FIELDS,
            0.2,
            0.2,
            4.083,
            2.941,
            32,
            0,
            breaks=False,
        )
        presentation.save(deck)

        MODULE.process(str(deck), floor_pt=10)

        checked = Presentation(deck)
        shape = next(s for s in checked.slides[0].shapes if s.has_text_frame)
        assert run_sizes(shape)[0] == 32


def test_paragraph_spacing_is_height_the_fitter_counts():
    """Space after a paragraph is real height no font size shrinks. Two boxes,
    identical but for a full inch of paragraph spacing, must not measure the
    same."""
    with tempfile.TemporaryDirectory() as temp_dir:
        deck = Path(temp_dir) / "spacing.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        plain = add_field_box(
            slide, "plain", PROMPT_FIELDS, 0.2, 0.2, 4.083, 2.941, 32, 0
        )
        spaced = add_field_box(
            slide, "spaced", PROMPT_FIELDS, 5.0, 0.2, 4.083, 2.941, 32, 72
        )
        presentation.save(deck)

        checked = Presentation(deck)
        shapes = {s.name: s for s in checked.slides[0].shapes if s.has_text_frame}
        plain_fit = MODULE.measure_shape(shapes["plain"], 32, 10)["best"]
        spaced_fit = MODULE.measure_shape(shapes["spaced"], 32, 10)["best"]
        assert spaced_fit < plain_fit


def test_a_line_break_element_counts_as_a_line():
    """pptxgenjs writes a newline character; hand-built and template text uses
    an <a:br/> element. Both start a line, so both are counted."""
    with tempfile.TemporaryDirectory() as temp_dir:
        deck = Path(temp_dir) / "br.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(0.2), Inches(0.2), Inches(4.0), Inches(2.0))
        frame = shape.text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        for index, field in enumerate(PROMPT_FIELDS):
            if index:
                paragraph.add_line_break()
            run = paragraph.add_run()
            run.text = field
            run.font.name = "Comic Sans MS"
            run.font.size = Pt(20)
        presentation.save(deck)

        checked = Presentation(deck)
        target = next(s for s in checked.slides[0].shapes if s.has_text_frame)
        lines = MODULE.paragraph_lines(
            target.text_frame.paragraphs[0], MODULE.FONT_REGULAR
        )
        assert len(lines) == len(PROMPT_FIELDS)
        assert [
            "".join(text for text, _ in line) for line in lines
        ] == PROMPT_FIELDS


if __name__ == "__main__":
    failed = 0
    for name, function in sorted(globals().items()):
        if name.startswith("test_") and callable(function):
            try:
                function()
                print("PASS", name)
            except AssertionError as error:
                failed += 1
                print("FAIL", name, str(error))
    raise SystemExit(1 if failed else 0)


def test_no_break_spaces_measure_as_one_unbroken_word():
    """The builder joins a calculation ("2,648 + 10 =") and the "Success
    Criteria" heading with U+00A0 so PowerPoint keeps them on one line. The
    measurer has to see the same thing, or it counts breaks the renderer
    never makes and leaves the box under-shrunk."""
    font = MODULE.pick_font_file(True, False)
    joined = "2,648 + 10 ="
    plain = "2,648 + 10 ="
    pt = 24
    whole_w, _ = MODULE._rendered_size(joined, pt, font)
    assert MODULE.widest_unbroken_word([(joined, font)], pt) >= whole_w * 0.99
    assert MODULE.widest_unbroken_word([(plain, font)], pt) < whole_w
    narrow = int(whole_w * 0.6)
    assert MODULE.wrap_paragraph(joined, narrow, pt, font) == 1
    assert MODULE.wrap_paragraph(plain, narrow, pt, font) >= 2


class _Box:
    """Just enough of a shape for the budget to measure: a size and some text."""

    class _Run:
        def __init__(self, text, size):
            self.text = text
            self.font = type("F", (), {"size": size, "bold": True, "italic": False})()

    class _Para:
        def __init__(self, text, size):
            self.runs = [_Box._Run(text, size)]
            self.text = text
            self.line_spacing = None
            self.space_before = None
            self.space_after = None

    class _Frame:
        def __init__(self, text, size):
            self.text = text
            self.paragraphs = [_Box._Para(text, size)]

    def __init__(self, text, width_in, height_in, pt=18):
        from pptx.util import Emu

        self.text_frame = _Box._Frame(text, Emu(int(pt * 12700)))
        self.width = Emu(int(width_in * 914400))
        self.height = Emu(int(height_in * 914400))
        self.name = "Text 1"


def test_budget_reports_volume_when_there_are_simply_too_many_words():
    """Too many words is answered by cutting or by a bigger zone, so the
    refusal gives the count to cut to."""
    box = _Box("Put the numbers in descending order and explain your reasoning fully.", 6.0, 0.4)
    message = MODULE.text_budget(box, 18, box.text_frame.text)
    assert "holds about" in message
    assert "this one is 69." in message


def test_budget_names_wrapping_when_the_words_fit_by_count():
    """A box with the area but not the width refuses text that is inside its
    character budget. Reporting only the budget reads as the build
    contradicting itself and sends the repair at the wording, which cannot
    help: the fix is a wider box."""
    box = _Box("Which column decides the order of the two numbers?", 1.2, 2.4)
    message = MODULE.text_budget(box, 18, box.text_frame.text)
    assert "wider" in message
    assert "holds about" not in message


def test_budget_says_nothing_rather_than_guessing_at_empty_text():
    box = _Box("   ", 3.0, 1.0)
    assert MODULE.text_budget(box, 18, box.text_frame.text) == ""
