"""A scheduled run knows which lesson of a long-term plan comes next, in any subject."""
from __future__ import annotations

import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

SCRIPT = Path(__file__).resolve().parents[1] / "plan-tracker.py"
SPEC = importlib.util.spec_from_file_location("plan_tracker", SCRIPT)
tracker = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(tracker)


def maths_plan(path: Path) -> Path:
    doc = Document()
    table = doc.add_table(rows=1, cols=8)
    for cell, text in zip(table.rows[0].cells, ["Unit", "National Curriculum Statement", "Lesson Objective", "Description", "Things to Look For", "Key Questions", "Sentence Stems", "#"]):
        cell.text = text
    rows = [
        ["Place Value 3 weeks", "Recognise place value", "LO: Count in 1,000s", "Count forwards", "", "", "", "1"],
        ["Place Value 3 weeks", "", "LO: Represent 4-digit numbers", "Use charts", "Zero as placeholder", "What is represented?", "There are ___ thousands", "2"],
        ["Addition 2 weeks", "Add numbers", "LO: Add two 4-digit numbers", "Column method", "", "", "", "3"],
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = text
    doc.save(str(path))
    return path


def kapow_plan(path: Path) -> Path:
    doc = Document()
    headers = ["Lesson from Kapow", "National Curriculum Statement", "Lesson Objective", "How Kapow would teach the lesson", "Key Vocabulary Kapow would use"]
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    def unit_row(title):
        cells = table.add_row().cells
        for cell in cells:
            cell.text = title
    def lesson_row(values):
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = text
    unit_row("Why are rainforests important to us?")
    lesson_row(["1. Where are rainforests?", "", "LO: To locate the Amazon", "Map work", "biome"])
    lesson_row(["2. Rainforest layers", "", "LO: To describe the layers", "Layer diagram", "canopy"])
    unit_row("Where does our food come from?")
    lesson_row(["1. Food choices", "", "LO: To explain food choices", "Discussion", "import"])
    doc.save(str(path))
    return path


class PlanTrackerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.plans = self.root / "letterbox"

    def tearDown(self):
        self.temp.cleanup()

    def run_cli(self, *args):
        import contextlib
        import io

        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            code = tracker.main(["--plans-dir", str(self.plans), *args])
        return code, out.getvalue()

    def next(self, plan):
        code, out = self.run_cli("next", "--plan", plan)
        self.assertEqual(code, 0, out)
        return json.loads(out)

    def test_a_maths_plan_with_a_unit_column_is_read_lesson_by_lesson(self):
        lessons = tracker.extract_lessons(maths_plan(self.root / "maths.docx"))
        self.assertEqual([(l["index"], l["unit"], l["lo"]) for l in lessons], [
            (1, "Place Value 3 weeks", "Count in 1,000s"),
            (2, "Place Value 3 weeks", "Represent 4-digit numbers"),
            (3, "Addition 2 weeks", "Add two 4-digit numbers"),
        ])
        self.assertIn("Things to Look For:\nZero as placeholder", lessons[1]["info"])
        self.assertIn("National Curriculum Statement:\nRecognise place value", lessons[1]["info"], "a statement written once carries to the rows below it")
        self.assertNotIn("#", lessons[0]["info"])

    def test_a_kapow_plan_with_units_spanning_the_row_is_read_too(self):
        lessons = tracker.extract_lessons(kapow_plan(self.root / "geography.docx"))
        self.assertEqual([(l["unit"], l["lo"]) for l in lessons], [
            ("Why are rainforests important to us?", "To locate the Amazon"),
            ("Why are rainforests important to us?", "To describe the layers"),
            ("Where does our food come from?", "To explain food choices"),
        ])
        self.assertIn("Lesson from Kapow:\n2. Rainforest layers", lessons[1]["info"])
        self.assertIn("Key Vocabulary Kapow would use:\ncanopy", lessons[1]["info"])

    def test_the_next_lesson_follows_the_last_built_and_stops_at_the_buffer(self):
        self.run_cli("import", str(maths_plan(self.root / "m.docx")), "--plan", "year4-maths", "--year", "Year 4", "--subject", "Maths", "--buffer", "2")
        first = self.next("year4-maths")
        self.assertEqual((first["status"], first["index"], first["lo"]), ("build", 1, "Count in 1,000s"))
        self.run_cli("advance", "--plan", "year4-maths", "--index", "1")
        self.run_cli("advance", "--plan", "year4-maths", "--index", "2")
        full = self.next("year4-maths")
        self.assertEqual((full["status"], full["reason"]), ("skip", "buffer_full"))
        self.run_cli("set-filed", "--plan", "year4-maths", "--index", "1")
        third = self.next("year4-maths")
        self.assertEqual((third["status"], third["index"]), ("build", 3))
        self.assertIn("opens a new unit", third["prior_context"])

    def test_a_finished_plan_does_nothing(self):
        self.run_cli("import", str(kapow_plan(self.root / "g.docx")), "--plan", "year4-geography", "--year", "4", "--subject", "Geography", "--start-after", "3")
        self.assertEqual(self.next("year4-geography")["reason"], "plan_complete")

    def test_only_the_next_lesson_can_be_marked_built(self):
        self.run_cli("import", str(maths_plan(self.root / "m.docx")), "--plan", "year4-maths", "--year", "4", "--subject", "Maths")
        code, out = self.run_cli("advance", "--plan", "year4-maths", "--index", "2")
        self.assertEqual(code, 1)
        self.assertIn("next lesson to mark built is 1", out)

    def test_saved_never_moves_backwards_and_the_teacher_can_move_the_plan(self):
        self.run_cli("import", str(maths_plan(self.root / "m.docx")), "--plan", "year4-maths", "--year", "4", "--subject", "Maths")
        self.run_cli("set-filed", "--plan", "year4-maths", "--index", "2")
        self.run_cli("set-filed", "--plan", "year4-maths", "--index", "1")
        self.assertEqual(self.next("year4-maths")["filed_up_to"], 2)
        code, out = self.run_cli("move", "--plan", "year4-maths", "--to", "1")
        self.assertIn("builds lesson 2 next (Represent 4-digit numbers)", out)
        after = self.next("year4-maths")
        self.assertEqual((after["built_up_to"], after["filed_up_to"], after["index"]), (1, 1, 2))

    def test_daily_and_weekly_subjects_get_different_default_buffers(self):
        self.run_cli("import", str(maths_plan(self.root / "m.docx")), "--plan", "year4-maths", "--year", "4", "--subject", "Maths")
        self.run_cli("import", str(kapow_plan(self.root / "g.docx")), "--plan", "year4-geography", "--year", "4", "--subject", "Geography")
        self.assertEqual(self.next("year4-maths")["buffer"], 5)
        self.assertEqual(self.next("year4-geography")["buffer"], 2)

    def test_built_and_saved_are_separate_files_so_two_computers_never_overwrite_each_other(self):
        self.run_cli("import", str(maths_plan(self.root / "m.docx")), "--plan", "year4-maths", "--year", "4", "--subject", "Maths")
        folder = self.plans / "plans" / "year4-maths"
        self.assertEqual(sorted(p.name for p in folder.iterdir()), ["built.json", "filed.json", "lessons.json", "plan.json"])

    def test_the_brief_carries_the_plan_row_for_the_designer(self):
        self.run_cli("import", str(maths_plan(self.root / "m.docx")), "--plan", "year4-maths", "--year", "4", "--subject", "Maths")
        brief = self.root / "brief.md"
        self.run_cli("next", "--plan", "year4-maths", "--brief", str(brief))
        text = brief.read_text(encoding="utf-8")
        self.assertIn("Year 4 Maths, lesson 1 of 3 in the long-term plan.", text)
        self.assertIn("Learning objective: Count in 1,000s", text)

    def test_a_plan_published_from_one_computer_is_opened_on_another(self):
        import subprocess

        def git(*args, cwd=None):
            done = subprocess.run(["git", *args], cwd=cwd, capture_output=True, text=True)
            self.assertEqual(done.returncode, 0, done.stderr)
            return done.stdout

        remote = self.root / "remote.git"
        git("init", "-q", "--bare", str(remote))
        teacher = self.root / "teacher"
        git("clone", "-q", str(remote), str(teacher))
        git("-c", "user.name=t", "-c", "user.email=t@t", "commit", "-q", "--allow-empty", "-m", "start", cwd=teacher)
        git("push", "-q", "origin", "HEAD:refs/heads/claude/lesson-outbox", cwd=teacher)

        self.plans = teacher
        self.assertEqual(self.run_cli("open")[0], 0)
        self.run_cli("import", str(maths_plan(self.root / "m.docx")), "--plan", "year4-maths", "--year", "4", "--subject", "Maths", "--start-after", "1")
        code, out = self.run_cli("publish", "--plan", "year4-maths")
        self.assertEqual(code, 0, out)

        cloud = self.root / "cloud"
        git("clone", "-q", str(remote), str(cloud))
        self.plans = cloud
        code, out = self.run_cli("open")
        self.assertEqual(code, 0, out)
        self.assertEqual(self.next("year4-maths")["index"], 2)
        self.run_cli("advance", "--plan", "year4-maths", "--index", "2")
        self.assertIn("published", self.run_cli("publish", "--plan", "year4-maths")[1])

        self.plans = teacher
        self.run_cli("open")
        self.assertEqual(self.next("year4-maths")["built_up_to"], 2, "the cloud's built counter reached the teacher's copy")

    def test_a_badly_named_plan_is_refused(self):
        code, out = self.run_cli("next", "--plan", "Year 4 Maths")
        self.assertEqual(code, 1)
        self.assertIn("lower-case words joined by hyphens", out)


if __name__ == "__main__":
    unittest.main()
