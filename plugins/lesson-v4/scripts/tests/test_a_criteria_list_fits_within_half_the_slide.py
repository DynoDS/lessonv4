"""A criteria list the panels the guidance names cannot hold is caught at design time.

A practice slide's success-criteria panel widens itself only as far as its list
needs to read at 18pt, and never past 6.35in; the half-width split is a roomier
option whose side is 0.15in taller, while shorter lists may use narrower free
geometry when measured fit permits it. Nobody after the lesson designer may reword a criterion, so the teacher
decided on 23 September 2026 that a list too long even for the widest box is
caught by the lesson check before any slides are made, and the lesson designer,
who owns the words, tightens it. The check measures exactly the shapes the
guidance names, so a list it passes is one the slide designer can place.

The lesson check measures in Python what the builder measures in JavaScript.
These tests hold the two to the same verdict: on every real list the long-list
investigation collected, on its six long lists in the teacher's style, and on
lists made either side of the new limit, including lists the widest practice
panel refuses and the half-width side holds, and lists only a wider band holds.

A list the lesson designer has tightened and still cannot fit is marked
`tooLongForPanels` and drawn smaller, down to 16pt, in the practice panel at
its widest or the half-width side, on a finished slide flagged for the teacher
(his ruling of 24 September 2026). The check lets a marked list through only
where the builder can draw it so, and these tests hold that verdict to the
builder's too, and the final text fit's floor for a marked list's lines.
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
import test_lesson_design_contract as contract  # noqa: E402

validator = contract.module
FIXTURE = json.loads((ROOT / "builder" / "test" / "fixtures" / "criteria-lists-4.2.288.json").read_text(encoding="utf-8"))
STAR = chr(0x2728)

# The builder's own answer for each list: the width the practice templates give
# it (or null when no width holds it at 18pt), and whether the panels of the
# roomiest free layouts hold it; and, for a list neither the practice panel nor
# the half-width side holds, the same two answers once the list is marked too
# long, when the builder draws it smaller, down to 16pt.
BUILDER = r"""
const B = process.argv[1];
console.warn = () => {};
const P = require(B + '/src/require-global')('pptxgenjs');
const { drawSlide } = require(B + '/src/templates');
const { scPanelWidth } = require(B + '/src/templates/maths-turn-sc');
const { listKey } = require(B + '/src/marked-criteria');
const lists = JSON.parse(require('fs').readFileSync(0, 'utf8'));
const ctx = (marked) => ({ slideIndex: 0, cardLook: true, imageDims: {}, markedCriteria: marked || new Set() });
const text = { type: 'text', value: 'Round 346 to the nearest ten.' };
const holds = (slide, marked) => {
  try { drawSlide(new P(), new P().addSlide(), slide, ctx(marked)); return true; }
  catch (err) { if (/^STEP_TEXT_OVERLOAD/.test(err.message)) return false; throw err; }
};
const widthFor = (steps, marked) => {
  try { return scPanelWidth({ criteria: { type: 'steps', steps } }, ctx(marked)); }
  catch (err) { if (!/^STEP_TEXT_OVERLOAD/.test(err.message)) throw err; return null; }
};
process.stdout.write(JSON.stringify(lists.map((steps) => {
  const panel = { type: 'sc-panel', content: { type: 'steps', steps } };
  const halfSlide = { template: 'split-h-50-50', title: 'T', primarySide: 'left', primary: text, secondary: panel };
  const verdict = {
    practice: widthFor(steps),
    half: holds(halfSlide),
    band: holds({ template: 'split-v-60-40', title: 'T', primarySide: 'top', primary: panel, secondary: text }),
    callouts: holds({ template: 'central-callouts-4', title: 'T', centre: panel, callouts: [text, text, text, text] }),
  };
  if (verdict.practice === null && !verdict.half) {
    const marked = new Set([listKey(steps)]);
    verdict.markedPractice = widthFor(steps, marked);
    verdict.markedHalf = holds(halfSlide, marked);
  }
  return verdict;
})));
"""

# The size of every panel the check measures, as the builder draws it.
ZONES = r"""
const B = process.argv[1];
console.warn = () => {};
const panel = require(B + '/src/success-criteria-panel');
const draw = panel.drawSuccessCriteriaPanel;
let zones = [];
panel.drawSuccessCriteriaPanel = (pptx, slide, zone, data, ctx) => { zones.push([zone.w, zone.h]); return draw(pptx, slide, zone, data, ctx); };
const P = require(B + '/src/require-global')('pptxgenjs');
const { drawSlide } = require(B + '/src/templates');
const long = 'Write each digit in its own column, lined up under the digit above it, and say its value.';
const steps = (n, s) => ({ type: 'steps', steps: Array.from({ length: n }, () => s) });
const text = { type: 'text', value: 'x' };
const out = {};
const record = (name, slide) => { zones = []; try { drawSlide(new P(), new P().addSlide(), slide, { slideIndex: 0, cardLook: true, imageDims: {} }); } catch (err) {} out[name] = zones; };
record('practice', { template: 'maths-turn-sc', title: 'T', questions: ['Q'], workingSpace: true, criteria: steps(9, long) });
record('half', { template: 'split-h-50-50', title: 'T', primarySide: 'left', primary: text, secondary: { type: 'sc-panel', content: steps(1, 'x') } });
record('band', { template: 'split-v-60-40', title: 'T', primarySide: 'top', primary: { type: 'sc-panel', content: steps(1, 'x') }, secondary: text });
record('callouts', { template: 'central-callouts-4', title: 'T', centre: { type: 'sc-panel', content: steps(1, 'x') }, callouts: [text, text, text, text] });
process.stdout.write(JSON.stringify(out));
"""


def run_builder(script: str, lists: list | None = None):
    builder = ROOT / "builder"
    out = subprocess.run(
        ["node", "-e", script, str(builder)], cwd=builder, input=json.dumps(lists or []),
        capture_output=True, text=True, encoding="utf-8", check=True,
    )
    return json.loads(out.stdout)


def made_lists() -> list[list[str]]:
    """Steps built from the real criteria's own words, one to nine of them, each
    from about 15 to 129 characters long in steps of two, some with sticky lines;
    then five to eight of them a character apart around the practice panel's
    limit, where the half-width side's extra 0.15in decides."""
    words = " ".join(" ".join(item["steps"]) for item in FIXTURE["lists"]).replace(STAR, "").split()
    made = []
    cursor = 0
    shapes = [(count, length) for count in range(1, 10) for length in range(15, 130, 2)]
    shapes += [(count, length) for count in range(5, 9) for length in range(44, 84)]
    for count, length in shapes:
        steps = []
        for _ in range(count):
            step = []
            while len(" ".join(step)) < length:
                step.append(words[cursor % len(words)])
                cursor += 1
            steps.append(" ".join(step))
        made.append(steps)
        if count in (4, 6) and length % 10 == 5:
            made.append(steps[:-1] + [f"{STAR} {steps[-1]}"])
            made.append(steps[:-2] + [f"{STAR} {steps[-2]}", f"{STAR} {steps[-1]}"])
    return made


CORPUS = [item["steps"] for item in FIXTURE["lists"]]
MADE = made_lists()
VERDICTS = run_builder(BUILDER, CORPUS + MADE)


def builder_holds(verdict: dict) -> bool:
    """Whether a shape the guidance names holds the list: the practice panel at
    any of its widths, or the half-width side."""
    return verdict["practice"] is not None or verdict["half"]


def test_the_lesson_check_and_the_builder_give_the_same_verdict():
    disagree = []
    for steps, verdict in zip(CORPUS + MADE, VERDICTS):
        fits = validator.criteria_fit(steps)["fits"]
        drawn = builder_holds(verdict)
        stars = sum(step.lstrip().startswith(STAR) for step in steps)
        # With two sticky lines the builder can be the stricter; the check may
        # then let a list through, and must never refuse one the builder draws.
        if fits != drawn and not (stars >= 2 and fits and not drawn):
            disagree.append((fits, verdict, steps))
    assert disagree == [], disagree[:5]


def test_either_side_of_the_new_limit():
    made = list(zip(MADE, VERDICTS[len(CORPUS):]))
    one_star = lambda steps: sum(step.startswith(STAR) for step in steps) < 2  # noqa: E731
    # Lists every practice width refuses and the half-width side holds: the
    # check lets each one through, because the guidance sends it there.
    half_side_only = [steps for steps, verdict in made
                      if verdict["practice"] is None and verdict["half"] and one_star(steps)]
    # Lists neither holds, though a wider band or a callout centre may: the
    # guidance never sends a list there, so the check refuses each one.
    band_only = [steps for steps, verdict in made
                 if not builder_holds(verdict) and (verdict["band"] or verdict["callouts"])]
    beyond_every_named_panel = [steps for steps, verdict in made if not builder_holds(verdict)]
    assert len(half_side_only) >= 5, len(half_side_only)
    assert len(band_only) >= 5, len(band_only)
    assert len(beyond_every_named_panel) >= 40, len(beyond_every_named_panel)
    for steps in half_side_only:
        fit = validator.criteria_fit(steps)
        assert fit["fits"] and fit["holders"] == ["the half-width side"], (fit["holders"], steps)
    assert not any(validator.criteria_fit(steps)["fits"] for steps in beyond_every_named_panel if one_star(steps))
    assert not any(validator.criteria_fit(steps)["fits"] for steps in band_only if one_star(steps))
    # And the made lists reach every practice width, so the widths are tested too.
    widths = [verdict["practice"] for _steps, verdict in made]
    assert all(widths.count(w) >= 10 for w in (4.6, 5.5, 6.35))


def test_a_marked_list_passes_exactly_where_the_builder_draws_it_smaller():
    # His ruling of 24 September 2026: a list the lesson designer marked too
    # long is drawn at the largest size that fits, down to 16pt, in the
    # practice panel at its widest or the half-width side, on a finished slide
    # flagged for the teacher, never left as a blank page. The check's verdict
    # is the builder's: a marked list it finds fits smaller is drawn smaller,
    # and one it finds too long even at 16pt (passed with a note) is one the
    # builder cannot draw. A sticky line keeps 18pt, and the mark applies only
    # to a list whose own steps no named shape holds at 18pt, so a made list
    # whose steps fit at 18pt beside a sticky line that does not is left out:
    # marked or not, that sticky line is the slide designer's to move, and a
    # design, which keeps its sticky lines apart, never carries one.
    own = lambda steps: [step for step in steps if not step.lstrip().startswith(STAR)]  # noqa: E731
    beyond = [(steps, verdict) for steps, verdict in zip(CORPUS + MADE, VERDICTS)
              if not builder_holds(verdict) and sum(step.lstrip().startswith(STAR) for step in steps) < 2
              and not validator.criteria_fit(own(steps))["fits"]]
    drawn = lambda verdict: verdict["markedPractice"] is not None or verdict["markedHalf"]  # noqa: E731
    disagree = [(steps, verdict) for steps, verdict in beyond
                if validator.criteria_fit_smaller(steps)["fits"] != drawn(verdict)]
    assert disagree == [], disagree[:3]
    drawn_smaller = [steps for steps, verdict in beyond if drawn(verdict)]
    assert len(drawn_smaller) >= 20 and len(beyond) - len(drawn_smaller) >= 20, (len(drawn_smaller), len(beyond))
    # Some only the half-width side holds, so its 0.15in is measured at 16pt too.
    assert sum(1 for _steps, verdict in beyond if verdict["markedPractice"] is None and verdict["markedHalf"]) >= 3
    # The practice panel widens before it draws a list smaller: a marked list
    # is drawn under 18pt only at 6.35in.
    assert {verdict["markedPractice"] for _steps, verdict in beyond} <= {None, 6.35}


def test_the_check_measures_the_shapes_the_guidance_names_at_the_builders_sizes():
    zones = run_builder(ZONES)
    practice_widths = sorted({w for w, _h in zones["practice"]})
    assert practice_widths == [4.6, 5.5, 6.35]
    named = [(w, 6.5) for w in practice_widths] + [tuple(zones["half"][0])]
    assert [(w, h) for _name, w, h in validator.NAMED_PANELS] == named
    # The wider band and the callout centre are not among them.
    measured = {(w, h) for _name, w, h in validator.NAMED_PANELS}
    assert tuple(zones["band"][0]) not in measured and tuple(zones["callouts"][0]) not in measured


def test_every_real_list_and_every_long_list_in_his_style_passes():
    refused = [item["source"] for item in FIXTURE["lists"] if not validator.criteria_fit(item["steps"])["fits"]]
    assert refused == []


# A list in his style that no named panel holds at 18pt and the practice panel
# at its widest holds at 17pt; and one that only the half-width side holds, at
# 16pt.
SMALLER = [
    "Round each number to the nearest hundred before you add, and write them down.",
    "Add the rounded numbers in your head and write the estimate beside the question.",
    "Work out the exact answer with the column method, lining up every column.",
    "Compare the exact answer with your estimate and say whether they are close.",
    "If they are not close, check each column again, starting with the ones.",
    "Write a sentence that says how your estimate helped you check the answer.",
    "Circle the answer you trust most and explain to a partner why you trust it.",
]
HALF_SIDE_ONLY = [
    "Write the 3-digit number on top and the 1-digit number under the ones column.",
    "Multiply the ones first, and say the multiplication out loud as you do it.",
    "If the answer is 10 or more, carry the tens under the next column to the left.",
    "Multiply the tens, then add any tens you carried before you write anything.",
    "If that is 10 or more, carry the hundreds under the hundreds column the same way.",
    "Multiply the hundreds and add anything you carried, then write the answer.",
    "Read the whole answer back, from the hundreds to the ones, saying each value.",
    "Check your answer with an estimate, rounding to the nearest hundred first.",
]


def smaller_design(steps=SMALLER):
    design, photos = contract.valid_contract()
    design["successCriteria"][0]["content"]["steps"] = list(steps)
    return design, photos


def too_long_design():
    design, photos = contract.valid_contract()
    design["successCriteria"][0]["content"]["steps"] = [
        "Write the 3-digit number on top and the 1-digit number under the ones column, lining up every digit.",
        "Multiply the ones first, and say the multiplication out loud as you do it so a partner can check it.",
        "If the answer is 10 or more, write the ones and carry the tens under the next column to the left.",
        "Multiply the tens, then add any tens you carried from the ones column before you write anything.",
        "If that is 10 or more, carry the hundreds under the hundreds column the same way as before.",
        "Multiply the hundreds and add anything you carried from the tens column, then write the answer.",
        "Read the whole answer back from the hundreds column to the ones column, saying each digit's value.",
        "Check your answer with an estimate, rounding the 3-digit number to the nearest hundred first.",
    ]
    return design, photos


def test_a_list_the_named_panels_cannot_hold_goes_back_to_its_writer():
    design, photos = too_long_design()
    try:
        validator.validate_design(design, photos)
    except validator.ContractError as exc:
        message = str(exc)
    else:
        raise AssertionError("an eight-step list of very long steps passed the lesson check")
    assert ("successCriteria[0] (sc-001), the list that begins \"Write the 3-digit number on top...\", "
            "is too long for the criteria panels slides are built with") in message
    assert "at 18pt, the smallest the board allows, its 8 steps take" in message
    assert "in the roomiest of them, the half-width side (6.35in wide and 6.65in tall), which holds" in message
    assert "and the practice panel holds it at none of its three widths." in message
    assert "Tighten the wording here, where it is written, keeping what each step tells a stuck child to do" in message
    assert "the slide designer has no roomier panel to give it, and nobody after you may reword a criterion." in message
    # This one is too long even drawn smaller, and the refusal says so; one
    # that fits at 16pt is not told that.
    assert ("Even at 16pt it does not fit: marked, it would pass with a note, but every slide that shows "
            "it would reach the teacher as a page to check with nothing drawn on it, so tighten it at least "
            "until it fits at 16pt.") in message
    assert "Even at 16pt" not in refusal(*smaller_design())
    # Decisions 3, 12 and 13 of the success-criteria topic: never fewer
    # criteria, and nothing downstream is told to shorten or drop one.
    for word in ("fewer", "drop", "remove", "omit", "split", "shorten", chr(0x2014), chr(0x2013)):
        assert word not in message.replace("split-h", ""), word


def test_the_numbers_in_the_message_show_why():
    design, _photos = too_long_design()
    steps = design["successCriteria"][0]["content"]["steps"]
    fit = validator.criteria_fit(steps)
    assert not fit["fits"] and fit["holders"] == []
    assert sum(fit["roomiest"]["lines"]) > fit["roomiest"]["holds"]
    assert 30 <= fit["roomiest"]["chars"] <= 45


def test_a_list_that_fits_passes_and_the_contract_list_passes():
    design, photos = contract.valid_contract()
    validator.validate_design(design, photos)
    design["successCriteria"][0]["content"]["steps"] = [
        "Write both numbers in a place value chart, one under the other.",
        "Compare the thousands digits first.",
        "If they are the same, compare the hundreds.",
        "Keep moving right until two digits are different.",
        "The number with the greater digit is the greater number.",
        "Put < or > between the numbers, with the open side facing the greater number.",
    ]
    validator.validate_design(design, photos)


def test_the_letter_widths_are_the_builders():
    table, unknown, safety = validator.comic_bold_widths()
    assert table["W"] == 1.0396 and table["l"] == 0.2739
    assert unknown == 1.1 and safety == 1.03


def test_the_guidance_says_what_the_builder_and_the_check_now_do():
    ssc = " ".join((ROOT / "references" / "slide-success-criteria.md").read_text(encoding="utf-8").split())
    # Delivery check for the owning guidance, not evidence of visual judgement.
    # Preserve engine thresholds and the source-owner exception while retiring
    # the blanket half-width fallback and categorical narrow-panel prohibition.
    assert "18pt minimum through both initial layout and the final text-fitting pass" in ssc
    assert "4.60in, then 5.50in, then 6.35in, never beyond half the slide" in ssc
    assert "half-area limit is a ceiling, not an allocation target" in ssc
    assert "A narrow column, a band or a panel sharing a column is suitable when" in ssc
    assert "all its content and furniture remain readable and the task stays usable" in ssc
    assert "Only that owner may tighten wording" in ssc
    assert "`tooLongForPanels`" in ssc and "down to 16pt" in ssc
    assert "a sticky line beside it stays at 18pt" in ssc
    assert "`CRITERIA_BELOW_READABLE_FLOOR`" in ssc
    assert "delivered flagged for the teacher to check, never cut to fit" in ssc
    assert "Never put a method's steps in a 30% column" not in ssc
    assert "or a slide needs a free layout, use the half-width split" not in ssc
    assert "as much as the widest practice panel holds" not in ssc
    templates = " ".join((ROOT / "references" / "templates.md").read_text(encoding="utf-8").split())
    assert "**The panel widens itself for a long list.** Across the `*-sc` family" in templates
    assert "the builder puts it above the working space instead." in templates


def refusal(design, photos) -> str:
    try:
        validator.validate_design(design, photos)
    except validator.ContractError as exc:
        return str(exc)
    raise AssertionError("the lesson check passed a list the named panels cannot hold")


def marked(design, index: int = 0):
    design["successCriteria"][index]["tooLongForPanels"] = True
    return design


def test_a_list_it_cannot_tighten_is_drawn_smaller_never_left_blank_and_the_refusal_says_what_that_costs():
    # Decision 11 of the success-criteria topic, in his words: "There is not a
    # time where I want the PowerPoint slide deck to never be produced because
    # of an error." A design the validator still refuses after its repair
    # passes ends the run with no deck, so the refusal says what to do when
    # tightening has failed, and what that costs. His ruling of 24 September
    # 2026: not a blank page but the list drawn smaller, down to 16pt, on a
    # finished slide flagged for him (the builder test "a list marked too long
    # is drawn smaller on a finished slide").
    design, photos = smaller_design()
    message = refusal(design, photos)
    assert ("Only if you have tightened it and no step can lose a word without losing what it tells a "
            "stuck child to do, keep the words, mark the list `\"tooLongForPanels\": true`, and say why in "
            "`flagsForTeacher`, in plain words for the teacher that name the list by its words.") in message
    assert "The check reads only the mark." in message
    assert ("It then lets the list through, and it costs every slide that shows the list: each draws it "
            "smaller than the 18pt the board allows everything else, down to 16pt, close to the smallest a "
            "class can read from the back of the room, and is flagged for the teacher to check before "
            "teaching; a list too long even at 16pt cannot be drawn at all, and each of its slides reaches "
            "the teacher as a page to check with nothing drawn on it. So tighten it if any word can go.") in message
    # Asked to tighten first.
    assert message.index("Tighten the wording here") < message.index("Only if you have tightened it")
    validator.validate_design(marked(design), photos)
    # A list only the half-width side holds at 16pt passes marked too.
    assert validator.criteria_fit_smaller(HALF_SIDE_ONLY)["holder"] == "the half-width side"
    design, photos = smaller_design(HALF_SIDE_ONLY)
    refusal(design, photos)
    validator.validate_design(marked(design), photos)


def test_a_marked_list_too_long_even_at_16pt_passes_with_a_note_and_the_deck_is_still_made():
    # Decision 11: "There is not a time where I want the PowerPoint slide deck to
    # never be produced because of an error." Drawn at 16pt the list would
    # still not fit, but the check lets it through with a plain note rather
    # than end the run with no deck; the review page asks for it to be
    # tightened; and the build writes the deck, each slide that shows the list
    # a page to check, the rarest case.
    import tempfile
    design, photos = too_long_design()
    validator.validate_design(marked(design), photos)
    notes = validator.criteria_marker_notes(design)
    assert len(notes) == 1
    assert notes[0].startswith("successCriteria[0] (sc-001), the list that begins \"Write the 3-digit number on "
                               "top...\", is marked `tooLongForPanels`, but it is too long even at 16pt, the least "
                               "a marked list is drawn at:")
    assert "in the roomiest criteria panel, the half-width side, which holds" in notes[0]
    assert ("Every slide that shows it will reach the teacher as a page to check before teaching, with its "
            "question, working space and criteria not drawn. Tighten it here, keeping what each step tells a "
            "stuck child to do, at least until it fits at 16pt, and better until it fits at 18pt and needs no "
            "mark; the design reviewer is asked to send it back to you. This is a note; the check passed, "
            "because the deck is always made.") in notes[0]
    for word in ("fewer", "drop", "remove", "omit", "shorten", chr(0x2014), chr(0x2013)):
        assert word not in notes[0], word
    # Its numbers are the 16pt measure's, and show the fault.
    smaller = validator.criteria_fit_smaller(design["successCriteria"][0]["content"]["steps"])
    assert not smaller["fits"] and sum(smaller["roomiest"]["lines"]) > smaller["roomiest"]["holds"]
    assert f"which holds {smaller['roomiest']['holds']} lines of about {smaller['roomiest']['chars']} characters" in notes[0]
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        (work / "lesson-design.json").write_text(json.dumps(design), encoding="utf-8")
        (work / "photo-requirements.json").write_text(json.dumps(photos), encoding="utf-8")
        # The check, as the run calls it: it passes, with the note first.
        out = subprocess.run([sys.executable, "-X", "utf8", str(ROOT / "scripts" / "validate-lesson-design.py"),
                              str(work / "lesson-design.json"), str(work / "photo-requirements.json")],
                             capture_output=True, text=True, encoding="utf-8")
        assert out.returncode == 0, out.stderr
        assert out.stdout.strip() == "LESSON_DESIGN_OK"
        assert out.stderr.splitlines()[0].startswith("LESSON_DESIGN_NOTE: successCriteria[0] (sc-001)")
        # Printed before the pass line, not only first on its own stream: run
        # unbuffered with both streams in one, in the order they were written.
        merged = subprocess.run([sys.executable, "-X", "utf8", "-u", str(ROOT / "scripts" / "validate-lesson-design.py"),
                                 str(work / "lesson-design.json"), str(work / "photo-requirements.json")],
                                stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8").stdout.splitlines()
        assert merged[0].startswith("LESSON_DESIGN_NOTE: ") and merged[-1] == "LESSON_DESIGN_OK", merged
        # And the deck: a practice slide showing the list, built as the run
        # builds it, is written with that slide a page to check.
        steps = design["successCriteria"][0]["content"]["steps"]
        (work / "lesson.json").write_text(json.dumps({
            "lessonName": "Too Long Even Smaller", "subject": "Maths", "lo": "Multiply a 3-digit number",
            "slides": [
                {"template": "split-v-60-40", "title": "Warm up", "primary": {"type": "text", "value": "Work out 4 x 7."},
                 "secondary": {"type": "text", "value": "Say it to your partner."}},
                {"template": "maths-turn-sc", "title": "My Turn", "questions": ["Work out 346 x 7."],
                 "workingSpace": True, "criteria": {"type": "steps", "steps": steps},
                 "speakerNotes": "Say to children: watch me."},
            ],
        }), encoding="utf-8")
        build = subprocess.run(["node", str(ROOT / "builder" / "build.js"), str(work / "lesson.json"), str(work / "out"),
                                "--deliver-flagged"], capture_output=True, text=True, encoding="utf-8")
        output = build.stdout + build.stderr
        if "AUTOFIT_DEPENDENCY_MISSING" in output or "AUTOFIT_NOT_PERMITTED" in output:
            return
        assert build.returncode == 0, output[-1500:]
        assert (work / "out" / "Too Long Even Smaller.pptx").exists()
        flagged = json.loads(next(line for line in build.stdout.splitlines()
                                  if line.startswith("SLIDES_FLAGGED: "))[len("SLIDES_FLAGGED: "):])
        assert flagged["slides"] == [2]
        diagnostics = [json.loads(line[len("BUILD_DIAGNOSTIC: "):]) for line in build.stdout.splitlines()
                       if line.startswith("BUILD_DIAGNOSTIC: ")]
        refused = [d for d in diagnostics if d["signal"] == "STEP_TEXT_OVERLOAD"]
        assert [d["location"]["slide"] for d in refused] == [2] and refused[0]["faultClass"] == "content"
        assert refused[0]["message"].startswith("This success-criteria list is marked too long in the design")


def test_the_check_reads_the_mark_and_never_the_words_of_a_flag():
    # The fourth check: reading a flag's words let a flag written for something
    # else through, and covered lists it did not name. No flag, however it is
    # worded, lets a list through now, and the mark needs no flag to be read.
    for flag in (
        "The plan's success criterion \"I can multiply any 3-digit number\" is too big for one lesson, "
        "so this lesson teaches 3-digit by 1-digit only.",
        "The supplied success criteria were too long for Year 4 to read, so I rewrote them.",
        "Write the 3-digit number on top ... is too long for the criteria panels: every step is needed.",
        "`sc-001 is too long for the criteria panels:` each step is one action a stuck child needs.",
    ):
        design, photos = too_long_design()
        design["flagsForTeacher"].append(flag)
        assert "successCriteria[0] (sc-001)" in refusal(design, photos), flag
    design, photos = smaller_design()
    validator.validate_design(marked(design), photos)
    # The mark is true or false, nothing else.
    design["successCriteria"][0]["tooLongForPanels"] = "yes"
    assert "successCriteria[0].tooLongForPanels" in refusal(design, photos)


def two_too_long_lists():
    design, photos = smaller_design()
    second = dict(design["successCriteria"][0], id="sc-002")
    second["content"] = {"steps": list(HALF_SIDE_ONLY)}
    design["successCriteria"].append(second)
    return design, photos


def test_two_lists_too_long_need_a_mark_each():
    design, photos = two_too_long_lists()
    message = refusal(design, photos)
    assert "successCriteria[0] (sc-001)" in message and "successCriteria[1] (sc-002)" in message
    assert "mark each list `\"tooLongForPanels\": true`" in message
    # Marking one lets only that one through.
    marked(design, 0)
    message = refusal(design, photos)
    assert "successCriteria[1] (sc-002)" in message and "successCriteria[0] (sc-001)" not in message
    validator.validate_design(marked(design, 1), photos)
    # Two lists that open alike are still two lists: marking one never covers
    # the other.
    design, photos = two_too_long_lists()
    design["successCriteria"][1]["content"]["steps"][0] = design["successCriteria"][0]["content"]["steps"][0]
    assert "successCriteria[1] (sc-002)" in refusal(marked(design, 0), photos)


def test_a_mark_on_a_list_the_check_does_not_find_too_long_is_a_note_not_a_refusal():
    design, photos = contract.valid_contract()
    validator.validate_design(marked(design), photos)
    assert validator.criteria_marker_notes(design) == [
        "successCriteria[0] (sc-001), the list that begins \"Partition each number.\", is marked "
        "`tooLongForPanels`, but it fits the criteria panels now: take the mark out, and the line in "
        "`flagsForTeacher` that explains it, so the teacher is not told something untrue. This is a "
        "note; the check passed."
    ]
    # A reference table is not measured, so a mark on one is named too.
    design["successCriteria"].append({
        "id": "sc-002", "type": "reference-table", "drawLive": False, "tooLongForPanels": True,
        "content": {"columns": ["Question says", "Operation"], "rows": [["altogether", "add"]]},
    })
    validator.validate_design(design, photos)
    assert "only a steps list is measured against the criteria panels" in validator.criteria_marker_notes(design)[1]
    # No flag's words ever draw a note: a flag that calls a plan's criterion
    # "too big", beside lists that all fit, is his to read, untouched.
    design, photos = contract.valid_contract()
    design["flagsForTeacher"].append(
        "The plan's success criterion \"I can add any two numbers\" is too big for one lesson.")
    validator.validate_design(design, photos)
    assert validator.criteria_marker_notes(design) == []


def load_packet():
    import importlib.util
    spec = importlib.util.spec_from_file_location("design_review_packet_fit", ROOT / "scripts" / "design-review-packet.py")
    packet = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(packet)
    return packet


def beside(view: str, sc_id: str) -> list[str]:
    """What the review page prints under one success-criteria list."""
    lines = view.splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith(f"- `{sc_id}` "))
    under = []
    for line in lines[start + 1:]:
        if not line.startswith("  - "):
            break
        under.append(line[4:])
    return under


def test_the_review_page_asks_beside_a_marked_list_whether_it_could_be_tightened_to_fit():
    # His ruling of 24 September 2026: a fix that changes what children read is
    # the lesson designer's, with the reviewer naming it. The page asks, on the
    # page and under the list itself, and does not say tightening was skipped.
    packet = load_packet()
    design, photos = smaller_design()
    under = beside(packet.build_review_view(marked(design), photos), "sc-001")
    asked = [line for line in under if line.startswith("Lesson check (read this one): ")]
    assert len(asked) == 1, under
    assert ("the lesson designer has marked it too long for them after trying to tighten it, so every "
            "slide that shows it will draw the list smaller than the 18pt floor, down to 16pt, and be "
            "flagged for the teacher to check before teaching.") in asked[0]
    assert ("Could the list be tightened until it fits at 18pt, with every step still telling a stuck "
            "child what to do? If so, return `REDESIGN REQUIRED` and name the tightening: the words are "
            "the lesson designer's to change.") in asked[0]
    assert "instead of tightening" not in asked[0] and "blank page" not in asked[0]
    # Beside a marked list too long even at 16pt, which passes the check with a
    # note, the reviewer is asked to send it back, naming the tightening.
    beyond, photos = too_long_design()
    asked = [line for line in beside(packet.build_review_view(marked(beyond), photos), "sc-001")
             if line.startswith("Lesson check (read this one): ")]
    assert len(asked) == 1
    assert ("no criteria panel a slide is built with holds this list even at 16pt, the least a list marked "
            "too long is drawn at, so every slide that shows it will reach the teacher as a page to check "
            "before teaching, with its question, working space and criteria not drawn. Return `REDESIGN "
            "REQUIRED` and name the tightening that brings it to 16pt at least, and to 18pt if it can, with "
            "every step still telling a stuck child what to do: the words are the lesson designer's to "
            "change.") in asked[0]
    # Beside a marked list that fits, it says so; beside an unmarked one, nothing.
    fits, photos = contract.valid_contract()
    assert any("marked too long for the criteria panels, but the check does not find it too long" in line
               for line in beside(packet.build_review_view(marked(fits), photos), "sc-001"))
    plain, photos = contract.valid_contract()
    assert not any(line.startswith("Lesson check") for line in beside(packet.build_review_view(plain, photos), "sc-001"))


def test_the_designer_contract_and_the_scaffold_carry_the_mark():
    template = " ".join((ROOT / "references" / "output-template.md").read_text(encoding="utf-8").split())
    assert ("`tooLongForPanels` is false (or left out) unless the lesson check refuses a steps list as too "
            "long for every criteria panel and you have tightened it and still cannot fit it without losing "
            "what a step tells a stuck child to do. Only then set it `true` on that list, and say why in "
            "`flagsForTeacher`. The check reads the mark and lets the list through, and every slide that "
            "shows it then draws it smaller than the 18pt floor, down to 16pt, flagged for the teacher to "
            "check before teaching, so it is never a way round tightening. A marked list too "
            "long even at 16pt still passes, with a note, because the deck is always made, but every slide "
            "that shows it reaches the teacher as a page to check with nothing drawn: tighten it at least "
            "that far.") in template
    # The reason's line sits in the list of what a flag is for.
    only_for = template[template.index("Put a concise string here only for:"):
                        template.index("Do not put ordinary design rationale here;")]
    assert ("- as a last resort, why a success-criteria list you marked `tooLongForPanels` could not be "
            "tightened, naming the list in plain words by its words.") in only_for
    designer = " ".join((ROOT / "agents" / "lesson-designer.md").read_text(encoding="utf-8").split())
    assert ("Mark that list `tooLongForPanels: true` as well: the check reads the mark, not the flag. Every "
            "slide that shows that list then draws it smaller than the 18pt floor, down to 16pt, flagged for "
            "the teacher to check, and one too long even at 16pt reaches the teacher as pages to check with "
            "nothing drawn, so tighten first.") in designer
    import test_lesson_design_scaffold as scaffold_tests
    generated, _photos = scaffold_tests.scaffold.build_scaffold(scaffold_tests.base_request())
    assert [item["tooLongForPanels"] for item in generated["successCriteria"]] == [False]


def test_nothing_offers_to_set_the_panel_width_and_the_widening_paragraph_ends_its_section():
    # The pin holds the template paragraph word for word; a paragraph written
    # after it could still undo it ("set panelWidth yourself"). There is no
    # such setting, and the paragraph is the last word of its section.
    templates = (ROOT / "references" / "templates.md").read_text(encoding="utf-8")
    for rel in ("references/templates.md", "references/slide-success-criteria.md", "agents/slide-designer.md",
                "builder/src/templates/maths-turn-sc.js"):
        assert "panelWidth" not in (ROOT / rel).read_text(encoding="utf-8").replace("panelWidening", ""), rel
    section = templates[templates.index("#### `maths-turn-sc`"):templates.index("#### `maths-turn-ref-sc`")]
    paragraphs = [" ".join(p.split()) for p in section.split("\n\n") if p.strip()]
    assert paragraphs[-1].startswith("**The panel widens itself for a long list.**"), paragraphs[-1][:80]


def test_the_final_text_fit_lets_only_a_marked_lists_lines_under_18pt():
    # The last check on the deck measures every box with the real font and
    # refuses one that needs less than its floor. A marked list's lines are
    # named for it by the builder, and only they may go under 18pt, never
    # under 16pt.
    import importlib.util
    import tempfile
    spec = importlib.util.spec_from_file_location("fit_text_postprocess_fit", ROOT / "builder" / "scripts" / "fit_text_postprocess.py")
    fit = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(fit)
    floor = lambda name: fit.shape_floor(name, fit.DEFAULT_FLOOR_PT)  # noqa: E731
    assert floor("GROWFIT__step-text-1-2-3-4__36__MIN16__marked-step-text-0") == 16
    # A sticky line beside a marked list keeps 18pt: only the list's own step
    # lines may go under it (the fifth check).
    assert floor("GROWFIT__step-reference-1-2-3-4__36__MIN17__marked-step-reference-2") == 18
    assert floor("GROWFIT__step-text-1-2-3-4__36__MIN12__marked-step-text-0") == 16
    assert floor("GROWFIT__step-text-1-2-3-4__36__MIN16__step-text-0") == 18
    assert floor("GROWFIT__step-text-1-2-3-4__36__MIN18__step-text-0") == 18
    assert floor("GROWFIT__table-1-2-3-4__40__MIN20__table-cell-0-0") == 20
    assert floor("GROWFIT__step-text-1-2-3-4__36__step-reference-0") == 18
    # The same words in the same box: at 16pt they fit, at 18pt they do not,
    # so the fit refuses the box unless it is a marked list's line.
    from pptx import Presentation
    from pptx.util import Inches, Pt
    words = "Add the rounded numbers in your head and write the estimate beside the question."
    with tempfile.TemporaryDirectory() as tmp:
        deck = Path(tmp) / "marked.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for top, name in ((0.5, "GROWFIT__g-1-2-3-4__36__MIN16__marked-step-text-0"),
                          (2.5, "GROWFIT__g-5-6-7-8__36__MIN18__step-text-0")):
            box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(5.0), Inches(0.75))
            box.name = name
            frame = box.text_frame
            frame.word_wrap = True
            frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
            run = frame.paragraphs[0].add_run()
            run.text, run.font.name, run.font.bold, run.font.size = words, "Comic Sans MS", True, Pt(16)
        prs.save(deck)
        out = subprocess.run([sys.executable, "-X", "utf8", str(ROOT / "builder" / "scripts" / "fit_text_postprocess.py"),
                              str(deck)], capture_output=True, text=True, encoding="utf-8")
        sizes = {shape.name.split("__")[-1]: shape.text_frame.paragraphs[0].runs[0].font.size.pt
                 for shape in Presentation(str(deck)).slides[0].shapes}
    overloaded = [line for line in out.stderr.splitlines() if "OVERLOAD" in line]
    assert len(overloaded) == 1 and "step-text-0" in overloaded[0] and "marked" not in overloaded[0], out.stderr
    assert 16 <= sizes["marked-step-text-0"] < 18, sizes


def test_a_mark_left_false_is_no_mark():
    # The scaffold writes `"tooLongForPanels": false` on every list, so a mark
    # left false must count for nothing: a list too long is refused as
    # unmarked, and a list that fits draws no note (the fifth check's V2).
    design, photos = too_long_design()
    design["successCriteria"][0]["tooLongForPanels"] = False
    message = refusal(design, photos)
    assert "successCriteria[0] (sc-001)" in message and "Only if you have tightened it" in message
    assert validator.criteria_marker_notes(design) == []
    packet = load_packet()
    assert not any(line.startswith("Lesson check") for line in beside(packet.build_review_view(design, photos), "sc-001"))
    design, photos = contract.valid_contract()
    design["successCriteria"][0]["tooLongForPanels"] = False
    validator.validate_design(design, photos)
    assert validator.criteria_marker_notes(design) == []


def test_the_note_reaches_whoever_runs_the_check_and_leaves_the_pass_exactly_as_it_was():
    import tempfile
    design, photos = contract.valid_contract()
    marked(design)
    with tempfile.TemporaryDirectory() as tmp:
        design_path = Path(tmp) / "lesson-design.json"
        photos_path = Path(tmp) / "photo-requirements.json"
        design_path.write_text(json.dumps(design), encoding="utf-8")
        photos_path.write_text(json.dumps(photos), encoding="utf-8")
        out = subprocess.run([sys.executable, "-X", "utf8", str(ROOT / "scripts" / "validate-lesson-design.py"),
                              str(design_path), str(photos_path)], capture_output=True, text=True, encoding="utf-8")
    assert out.returncode == 0, out.stderr
    assert out.stdout.strip() == "LESSON_DESIGN_OK"
    assert "LESSON_DESIGN_NOTE: successCriteria[0] (sc-001), the list that begins" in out.stderr
